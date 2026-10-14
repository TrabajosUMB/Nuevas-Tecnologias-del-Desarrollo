"""Extrae de la guía Lab3_Solidity_Contratos.docx el contrato de ejemplo tal como está escrito y lo guarda como MiContrato.sol.

Se extrae por programa (no se copia a mano) para que el contrato probado en Remix sea exactamente el de la guía.
Uso: py extraer_contrato.py
"""
from __future__ import annotations

import pathlib
import sys

import docx

sys.stdout.reconfigure(encoding="utf-8")
AQUI = pathlib.Path(__file__).resolve().parent.parent
GUIA = AQUI / "Lab3_Solidity_Contratos.docx"
SALIDA = AQUI / "contratos" / "MiContrato.sol"


def buscar_codigo(d: docx.document.Document) -> str:
    for tabla in d.tables:
        for fila in tabla.rows:
            for celda in fila.cells:
                if "SPDX-License-Identifier" in celda.text:
                    return "\n".join(p.text for p in celda.paragraphs)
    raise SystemExit("No se encontró el contrato en la guía")


if __name__ == "__main__":
    codigo = buscar_codigo(docx.Document(str(GUIA))).replace(" ", " ")
    SALIDA.parent.mkdir(exist_ok=True)
    SALIDA.write_text(codigo.rstrip("\n") + "\n", encoding="utf-8", newline="\n")
    print(f"{SALIDA} ({len(codigo.splitlines())} líneas)")
    print(codigo)
