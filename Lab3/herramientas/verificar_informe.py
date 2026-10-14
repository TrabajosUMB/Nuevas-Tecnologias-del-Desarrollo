"""Verificación del informe del Lab 3: metadatos, texto problemático y páginas (docx + pdf). Repone el autor tras el guardado de Word."""
from __future__ import annotations

import pathlib
import re
import sys

import docx
import pymupdf

sys.stdout.reconfigure(encoding="utf-8")
AQUI = pathlib.Path(__file__).resolve().parent.parent
DOCX = AQUI / "Lab3_Solidity_Grupo6.docx"
PDF = DOCX.with_suffix(".pdf")
AUTORES = "Rodríguez Ángel, Santiago; Olaya Rojas, Santiago"
PATRONES = {
    "mojibake": r"Ã.|Â.|â€|�",
    "referencia sin resolver": r"\?\?|\{[TFN]:",
    "marca HTML suelta": r"</?(b|i|bi|br)\s*/?>",
    "campo sin actualizar": r"Actualice este campo",
    "marcador pendiente": r"\[(cifra|dato|pendiente|completar)[^\]]*\]|TODO|XXX",
}


def restablecer_autor() -> None:
    d = docx.Document(str(DOCX))
    cp = d.core_properties
    cp.author, cp.last_modified_by = AUTORES, AUTORES
    d.save(str(DOCX))
    cp = docx.Document(str(DOCX)).core_properties
    print(f"metadatos: autor={cp.author!r} | último en modificar={cp.last_modified_by!r} | título={cp.title!r} | comentarios={cp.comments!r}")


def revisar_texto() -> None:
    d = docx.Document(str(DOCX))
    textos = [p.text for p in d.paragraphs]
    for t in d.tables:
        textos += [c.text for fila in t.rows for c in fila.cells]
    completo = "\n".join(textos)
    pdf_texto = "\n".join(pagina.get_text() for pagina in pymupdf.open(str(PDF)))
    for nombre, patron in PATRONES.items():
        print(f"{nombre}: docx={len(re.findall(patron, completo))} pdf={len(re.findall(patron, pdf_texto))}")
    print(f"secciones: {len(d.sections)} | páginas del PDF: {len(pymupdf.open(str(PDF)))}")


if __name__ == "__main__":
    restablecer_autor()
    revisar_texto()
