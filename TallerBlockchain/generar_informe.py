from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
import os

logo_path = r'C:\Users\cuent\Documents\Universidad\2026-02\Soft Seguro\LabPractico\logo2.png'
output_path = r'C:\Users\cuent\Documents\Universidad\2026-02\Nuevas-Tecnologias-del-Desarrollo\TallerBlockchain\Informe_TallerBlockchain_Santiago.docx'

doc = Document()

for section in doc.sections:
    section.top_margin    = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin   = Inches(1)
    section.right_margin  = Inches(1)
    section.page_width    = Inches(8.5)
    section.page_height   = Inches(11)

style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(12)
style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
style.paragraph_format.space_before = Pt(0)
style.paragraph_format.space_after  = Pt(0)


def h1(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    r = p.add_run(text)
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    return p


def h2(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    r = p.add_run(text)
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    return p


def body(text, indent=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(0)
    if indent:
        p.paragraph_format.first_line_indent = Inches(0.5)
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    return p


def caption(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    r = p.add_run(text)
    r.italic = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)
    return p


def item(text):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Inches(0.5)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(0)
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    return p


def ref_entry(text):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.5)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(0)
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    return p


# ===================== PORTADA =====================
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
if os.path.exists(logo_path):
    p.add_run().add_picture(logo_path, width=Inches(2.0))

for label, bold, size in [
    ('Universidad Manuela Beltrán', True, 14),
    ('Programa de Ingeniería de Software', False, 12),
    ('Asignatura: Nuevas Tecnologías del Desarrollo', False, 12),
]:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    r = p.add_run(label)
    r.bold = bold
    r.font.name = 'Times New Roman'
    r.font.size = Pt(size)

doc.add_paragraph()
doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
r = p.add_run('Informe de Laboratorio: Inmutabilidad y Consenso en Blockchain')
r.bold = True
r.font.name = 'Times New Roman'
r.font.size = Pt(16)

doc.add_paragraph()
doc.add_paragraph()

for label, bold in [
    ('Presentado por:', False),
    ('Rodríguez Angel, Santiago', True),
    ('Olaya Rojas, Santiago', True),
]:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    r = p.add_run(label)
    r.bold = bold
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)

doc.add_paragraph()
doc.add_paragraph()
doc.add_paragraph()

for label in ['Bogotá, Colombia', '2026']:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    r = p.add_run(label)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)

doc.add_page_break()

# ===================== INTRODUCCIÓN =====================
h1('Introducción')
body(
    'La tecnología blockchain representa uno de los avances más significativos de la era digital, '
    'siendo la base técnica de la Web 3.0. A diferencia de las bases de datos tradicionales '
    'centralizadas, una blockchain es un registro distribuido e inmutable en el que cada entrada, '
    'denominada bloque, está vinculada criptográficamente con la anterior, formando una cadena que '
    'garantiza la integridad de los datos sin necesidad de una autoridad central de confianza.'
)
body(
    'El presente informe documenta los resultados del laboratorio práctico de simulación de '
    'blockchain, en el que se exploró la anatomía de un bloque, el funcionamiento del algoritmo de '
    'consenso Proof of Work (PoW), el efecto avalancha de la función SHA-256 y la razón por la '
    'cual alterar una transacción histórica en una red descentralizada resulta computacionalmente '
    'inviable.'
)
body(
    'Para llevar a cabo la actividad se utilizó el simulador visual Blockchain Demo, desarrollado '
    'por Anders Brownworth, que permite observar en tiempo real el comportamiento de los hashes, '
    'el proceso de minería y la propagación de errores de integridad a lo largo de la cadena.'
)

doc.add_page_break()

# ===================== OBJETIVO =====================
h1('1. Objetivo de la Actividad')
body(
    'Comprender de manera práctica el funcionamiento de una cadena de bloques, analizando la '
    'relación entre el Nonce, la Data y el Hash, y cómo estos elementos garantizan la integridad '
    'e inmutabilidad de la información almacenada en la red.'
)

# ===================== PARTE A =====================
doc.add_paragraph()
h1('2. Parte A: El Bloque Único')
h2('2.1 Procedimiento Realizado')
body(
    'En la sección Block del simulador, se ingresó el nombre completo «Santiago Rodríguez Angel» '
    'en el campo Data. Se observó que, a medida que se escribía cada carácter, el hash SHA-256 '
    'del bloque cambiaba completamente en tiempo real, aunque la modificación en los datos fuese '
    'mínima. El bloque permaneció con fondo rojo, indicando un estado inválido, hasta que se '
    'presionó el botón Mine.'
)
body(
    'Tras ejecutar el proceso de minería, el simulador encontró un valor de Nonce que produce un '
    'Hash cuya representación hexadecimal comienza con cuatro ceros (0000…). Al alcanzar dicha '
    'condición, el bloque cambió su color de rojo a verde, señalando que el bloque es válido '
    'según las reglas del protocolo.'
)

h2('2.2 Captura de Pantalla')
caption('[Figura 1: Bloque válido tras la minería — Hash inicia con "0000", estado verde y Nonce encontrado]')

h2('2.3 Respuesta a la Pregunta del Taller')
body('Pregunta: ¿Qué sucede con el color del bloque y qué relación tiene el Hash resultante con el botón de minado?')
body(
    'Cuando se ingresa información en el campo Data sin minar, el bloque aparece en color rojo '
    'porque el hash calculado no cumple el criterio de dificultad establecido (comenzar con 0000). '
    'Al hacer clic en el botón Mine, el algoritmo Proof of Work incrementa iterativamente el valor '
    'del Nonce — un número entero arbitrario — y recalcula el hash en cada intento hasta encontrar '
    'aquel que satisfaga la condición de dificultad. Una vez hallado, el bloque cambia a color verde '
    'y el Hash resultante comienza con cuatro ceros. Esto demuestra que la minería no consiste en '
    '"inventar" el hash, sino en un proceso de prueba y error computacionalmente costoso que '
    'garantiza el consenso y la seguridad de la red.'
)

# ===================== PARTE B =====================
doc.add_paragraph()
h1('3. Parte B: La Cadena (Blockchain)')
h2('3.1 Procedimiento Realizado')
body(
    'En la sección Blockchain del simulador, se visualizó una cadena de cinco bloques enlazados. '
    'En el campo Data del Bloque 1 se escribió "Transacción 001" y se ejecutó la minería. Luego, '
    'en el Bloque 4, se modificó un carácter de su contenido, observando el efecto inmediato en '
    'la cadena.'
)
body(
    'De manera inmediata, el Bloque 4 y el Bloque 5 se tornaron rojos, indicando que su integridad '
    'fue comprometida. Posteriormente, se procedió a minar nuevamente el Bloque 4, logrando que '
    'este volviera a color verde. Sin embargo, el Bloque 5 permaneció en rojo.'
)

h2('3.2 Capturas de Pantalla')
caption('[Figura 2: Cadena tras modificar el Bloque 4 — Bloques 4 y 5 en rojo (inválidos)]')
caption('[Figura 3: Cadena tras re-minar el Bloque 4 — Bloque 4 verde, Bloque 5 permanece rojo]')

h2('3.3 Reflexión: ¿Por Qué el Bloque 5 Sigue en Rojo?')
body(
    'El Bloque 5 permanece inválido (rojo) después de re-minar el Bloque 4 por el siguiente motivo: '
    'cada bloque contiene en su cabecera el hash del bloque inmediatamente anterior, campo '
    'denominado «Prev Hash». Cuando el Bloque 4 fue modificado, su hash cambió. Al minarlo '
    'nuevamente, se calculó un nuevo hash distinto al que tenía originalmente. El Bloque 5 fue '
    'construido con el hash antiguo del Bloque 4 como su «Prev Hash»; al variar el hash del '
    'Bloque 4 — incluso después de una minería válida —, el «Prev Hash» que el Bloque 5 registra '
    'ya no coincide con el nuevo hash del Bloque 4, por lo que el Bloque 5 resulta '
    'criptográficamente inválido.'
)
body(
    'Este mecanismo constituye la esencia de la inmutabilidad blockchain: para alterar cualquier '
    'bloque del pasado sin ser detectado, el atacante tendría que re-minar ese bloque y todos los '
    'bloques posteriores de la cadena, además de hacerlo más rápido que el resto de la red '
    'distribuida, lo que resulta computacionalmente inviable en redes con miles de nodos.'
)

# ===================== PARTE C =====================
doc.add_paragraph()
h1('4. Parte C: El Libro Contable (Tokens)')
h2('4.1 Procedimiento Realizado')
body('En la sección Tokens del simulador se crearon tres transferencias ficticias:')
item('1. De Darcy → Binance: 50 tokens')
item('2. De Alice → Bob: 25 tokens')
item('3. De Bob → Carol: 10 tokens')
body(
    'Una vez registradas y minadas las tres transferencias, se intentó modificar la primera '
    'transacción del bloque más antiguo (cambiar el monto de 50 a 200 tokens), observando '
    'el efecto en cadena.'
)

h2('4.2 Capturas de Pantalla')
caption('[Figura 4: Sección Tokens con tres transacciones minadas y validadas]')
caption('[Figura 5: Efecto en cadena al modificar una transacción histórica — bloques posteriores en rojo]')

h2('4.3 Análisis: ¿Por Qué es Prácticamente Imposible Hackear una Red con Miles de Nodos?')
body(
    'La simulación evidencia que alterar cualquier transacción registrada en un bloque pasado '
    'produce un efecto en cascada que invalida todos los bloques subsiguientes de la cadena. '
    'Para "corregir" esta inconsistencia, el atacante necesitaría:'
)
item('(a) Re-minar el bloque alterado, proceso que requiere una cantidad significativa de poder computacional (trabajo de Proof of Work).')
item('(b) Re-minar cada uno de los bloques que siguen al bloque modificado, ya que cada uno hereda el hash del anterior.')
item(
    '(c) Completar todo este proceso antes de que la red legítima agregue un nuevo bloque, '
    'lo cual implicaría superar el 51 % del poder de cómputo total de todos los nodos de la red.'
)
body(
    'En redes como Bitcoin, que cuenta con miles de nodos distribuidos en todo el mundo y una tasa '
    'de hash total de exahashes por segundo, superar el 51 % del poder de cómputo resulta económica '
    'y tecnológicamente inviable para cualquier atacante. Cada nodo de la red mantiene una copia '
    'completa de la cadena y rechaza de forma autónoma cualquier versión que no cumpla las reglas '
    'de consenso, eliminando así el punto único de fallo que existiría en un sistema centralizado.'
)

# ===================== DEFINICIÓN HASH =====================
doc.add_page_break()
h1('5. Definición Técnica: ¿Qué es un Hash y por Qué es Fundamental?')
body(
    'Un hash criptográfico es el resultado de aplicar una función matemática de un solo sentido '
    '(one-way function) a un conjunto de datos de entrada de longitud arbitraria, produciendo una '
    'cadena de longitud fija. En el contexto blockchain se utiliza el algoritmo SHA-256 (Secure '
    'Hash Algorithm 256 bits), que genera siempre una salida de 64 caracteres hexadecimales '
    '(256 bits).'
)
body('Las propiedades que hacen al hash fundamental para la seguridad son:')
item('1. Determinismo: la misma entrada siempre produce exactamente el mismo hash.')
item(
    '2. Efecto avalancha: un cambio mínimo en la entrada (un carácter, un bit) produce un hash '
    'completamente diferente. Esto garantiza que cualquier alteración en los datos sea detectada de inmediato.'
)
item('3. Irreversibilidad: es computacionalmente imposible obtener los datos de entrada a partir del hash (función de un solo sentido).')
item('4. Resistencia a colisiones: es extremadamente improbable que dos entradas distintas produzcan el mismo hash.')
item(
    '5. Velocidad de verificación: verificar si un hash es correcto es instantáneo, mientras que '
    'encontrar un hash válido (minería) es computacionalmente costoso.'
)
body(
    'En la cadena de bloques, el hash cumple tres roles simultáneos: (1) sirve como identificador '
    'único del bloque, (2) vincula cada bloque con el anterior mediante la referencia al hash previo, '
    'y (3) actúa como sello de integridad que garantiza que ningún dato del bloque ha sido modificado '
    'desde su minería.'
)

# ===================== CASO DE USO =====================
doc.add_paragraph()
h1('6. Caso de Uso: Trazabilidad de Medicamentos en el Sector Salud')
body(
    'Una de las aplicaciones más relevantes de blockchain fuera del ámbito de las criptomonedas es '
    'la trazabilidad de medicamentos en la cadena farmacéutica. La falsificación de medicamentos '
    'representa un problema global que cobra vidas anualmente y genera pérdidas económicas '
    'billonarias (Organización Mundial de la Salud, 2022).'
)
body(
    'Mediante una blockchain permisionada, cada etapa del ciclo de vida de un medicamento — '
    'fabricación, control de calidad, distribución, almacenamiento y dispensación — puede '
    'registrarse como una transacción inmutable. Cada actor de la cadena (laboratorio, distribuidor, '
    'farmacia, hospital) firma digitalmente su intervención y la registra en la red.'
)
body('Las ventajas clave son:')
item(
    '1. Inmutabilidad: ningún actor puede alterar retroactivamente los registros de origen, '
    'temperatura o fecha de vencimiento.'
)
item(
    '2. Transparencia auditada: reguladores sanitarios como el INVIMA (Colombia) o la FDA '
    '(EE. UU.) pueden auditar la cadena completa en tiempo real sin depender de informes de terceros.'
)
item(
    '3. Trazabilidad instantánea: ante un retiro del mercado, es posible identificar en minutos '
    'qué lotes específicos se distribuyeron y a qué puntos de venta llegaron.'
)
item(
    '4. Smart Contracts: contratos inteligentes pueden automatizar el pago entre actores una vez '
    'que la entrega es confirmada por el receptor, eliminando intermediarios financieros.'
)
body(
    'Esta aplicación ilustra cómo la blockchain trasciende las criptomonedas para convertirse en '
    'infraestructura de confianza digital aplicable a cualquier sector donde la autenticidad, el '
    'origen y la integridad de los datos sean críticos.'
)

# ===================== CONCLUSIONES =====================
doc.add_page_break()
h1('7. Conclusiones')
body(
    'El laboratorio permitió comprender de forma empírica por qué la tecnología blockchain es '
    'considerada prácticamente inalterable. La combinación de tres mecanismos — la función de '
    'hash SHA-256, el Nonce y el enlace criptográfico entre bloques — crea un sistema donde '
    'modificar cualquier dato histórico requiere un esfuerzo computacional proporcional a toda '
    'la potencia de la red, lo cual resulta inviable en escenarios reales.'
)
body(
    'El efecto avalancha del hash SHA-256 garantiza que hasta el más mínimo cambio en la '
    'información sea detectado de forma inmediata por todos los nodos de la red, haciendo '
    'innecesaria la presencia de una autoridad central que certifique la veracidad de las transacciones.'
)
body(
    'Más allá de las criptomonedas, blockchain representa un paradigma de confianza distribuida '
    'aplicable a sectores como salud, logística, gobierno, propiedad intelectual e identidad '
    'digital, donde la transparencia y la integridad de los datos son requisitos críticos. La '
    'comprensión de sus fundamentos técnicos es indispensable para cualquier profesional de la '
    'ingeniería de software que aspire a construir soluciones de la Web 3.0.'
)

# ===================== REFERENCIAS =====================
doc.add_page_break()
h1('Referencias')
ref_entry(
    'Brownworth, A. (2023). Blockchain Demo. Anders Brownworth. '
    'https://andersbrownworth.com/blockchain/'
)
ref_entry(
    'Nakamoto, S. (2008). Bitcoin: A peer-to-peer electronic cash system. '
    'Bitcoin.org. https://bitcoin.org/bitcoin.pdf'
)
ref_entry(
    'National Institute of Standards and Technology. (2015). Secure Hash Standard (SHS) '
    '(FIPS PUB 180-4). U.S. Department of Commerce. https://doi.org/10.6028/NIST.FIPS.180-4'
)
ref_entry(
    'Tapscott, D., & Tapscott, A. (2016). Blockchain revolution: How the technology behind '
    'bitcoin is changing money, business, and the world. Portfolio/Penguin.'
)
ref_entry(
    'World Health Organization. (2022). Falsified medical products. WHO. '
    'https://www.who.int/news-room/fact-sheets/detail/substandard-and-falsified-medical-products'
)
ref_entry(
    'Zheng, Z., Xie, S., Dai, H., Chen, X., & Wang, H. (2017). An overview of blockchain '
    'technology: Architecture, consensus, and future trends. 2017 IEEE International Congress '
    'on Big Data, 557-564. https://doi.org/10.1109/BigDataCongress.2017.85'
)

doc.save(output_path)
print('OK:', output_path)
